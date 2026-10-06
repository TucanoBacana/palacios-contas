"""Importa os dados da planilha (cardapio + lancamentos) para o banco Postgres (Supabase).

Precisa da variavel de ambiente DATABASE_URL configurada (veja .env.example).

Usos comuns:
  python import_xlsx.py planilha.xlsx --desde 2026-09-28 --dry-run   # so mostra o que faria
  python import_xlsx.py planilha.xlsx --desde 2026-09-28             # acrescenta so o que falta
  python import_xlsx.py planilha.xlsx --reset                        # apaga tudo e reimporta

Sem --reset nada e apagado: o cardapio e atualizado e so entram lancamentos que ainda nao
existem no banco (mesma data, pessoa, item, quantidade e valor), a partir de --desde.
O estoque nao e mexido: as vendas da planilha sao historico.
"""
import argparse
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime

import openpyxl

from db import get_connection, get_or_create_pessoa, init_db


def normaliza_status(valor):
    if not valor:
        return False
    return str(valor).strip().lower().startswith("pago")


def to_iso_date(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, str) and valor.strip():
        return valor.strip()
    return date.today().isoformat()


def chave_nome(texto):
    """Nome sem acento, espacos repetidos e caixa: 'Débora ' e 'DEBORA' sao a mesma pessoa."""
    t = unicodedata.normalize("NFD", str(texto or ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t.strip().upper())


def importar(caminho_xlsx, reset=False, min_row_pedidos=2, desde=None, dry_run=False,
             categoria_novos="mercadinho"):
    init_db()
    conn = get_connection()
    cur = conn.cursor()

    if reset:
        cur.execute("DELETE FROM pedidos; DELETE FROM produtos; DELETE FROM pessoas;")

    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws_cardapio = wb.worksheets[1]
    ws_pedidos = wb.worksheets[0]

    # --- cardapio ---
    atualizados, novos = [], []
    for row in ws_cardapio.iter_rows(min_row=2, values_only=True):
        numero_item, nome, preco = row[0], row[1], row[2]
        if numero_item is None or nome is None:
            continue
        nome, preco = re.sub(r"\s+", " ", str(nome).strip()), float(preco or 0)
        cur.execute("SELECT id, nome, preco FROM produtos WHERE numero_item = %s", (numero_item,))
        existente = cur.fetchone()
        if existente:
            if existente["nome"] != nome or existente["preco"] != preco:
                atualizados.append(f"{numero_item}: {existente['nome']!r} -> {nome!r}, R$ {preco:.2f}")
            cur.execute("UPDATE produtos SET nome = %s, preco = %s WHERE id = %s", (nome, preco, existente["id"]))
        else:
            novos.append(f"{numero_item}: {nome!r}, R$ {preco:.2f} ({categoria_novos})")
            cur.execute(
                "INSERT INTO produtos (numero_item, nome, preco, categoria) VALUES (%s, %s, %s, %s)",
                (numero_item, nome, preco, categoria_novos),
            )

    # --- lancamentos ---
    cur.execute("SELECT id, nome FROM pessoas")
    pessoas = {chave_nome(p["nome"]): p["id"] for p in cur.fetchall()}
    cur.execute(
        """SELECT p.data, pe.nome, pr.numero_item, p.quantidade, p.valor_total
           FROM pedidos p JOIN pessoas pe ON pe.id = p.pessoa_id JOIN produtos pr ON pr.id = p.produto_id"""
    )
    ja_existe = Counter((r["data"].isoformat(), chave_nome(r["nome"]), r["numero_item"],
                         float(r["quantidade"]), round(r["valor_total"], 2)) for r in cur.fetchall())

    importados, repetidos, pessoas_novas, total = 0, 0, [], 0.0
    for row in ws_pedidos.iter_rows(min_row=min_row_pedidos, values_only=True):
        (data_val, pessoa, num_item, produto_nome, qtd, valor_unit,
         valor_total, forma_pag, pago, data_pagto, obs) = (list(row) + [None] * 11)[:11]

        if not pessoa or not num_item:
            continue
        data_iso = to_iso_date(data_val)
        if desde and data_iso < desde:
            continue

        cur.execute("SELECT id, preco, custo FROM produtos WHERE numero_item = %s", (num_item,))
        prod_row = cur.fetchone()
        if not prod_row:
            print(f"  aviso: item {num_item!r} nao encontrado no cardapio, pulando linha ({pessoa!r})")
            continue

        vu = float(valor_unit) if isinstance(valor_unit, (int, float)) else prod_row["preco"]
        qtd_f = float(qtd) if isinstance(qtd, (int, float)) else 1.0
        vt = float(valor_total) if isinstance(valor_total, (int, float)) else vu * qtd_f

        chave = (data_iso, chave_nome(pessoa), int(num_item), qtd_f, round(vt, 2))
        if ja_existe[chave] > 0:
            ja_existe[chave] -= 1
            repetidos += 1
            continue

        nome_pessoa = re.sub(r"\s+", " ", str(pessoa).strip())
        pessoa_id = pessoas.get(chave_nome(pessoa))
        if pessoa_id is None:
            pessoa_id = get_or_create_pessoa(conn, nome_pessoa)
            pessoas[chave_nome(pessoa)] = pessoa_id
            pessoas_novas.append(nome_pessoa)

        pago_bool = normaliza_status(pago)
        custo = float(prod_row["custo"] or 0)
        cur.execute(
            """INSERT INTO pedidos
               (data, pessoa_id, produto_id, quantidade, valor_unitario, valor_total,
                custo_unitario, custo_total, forma_pagamento, pago, valor_pago, data_pagamento,
                observacoes, criado_por)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'Planilha')""",
            (
                data_iso, pessoa_id, prod_row["id"], qtd_f, vu, vt, custo, round(custo * qtd_f, 2),
                str(forma_pag).strip() if forma_pag else None,
                pago_bool, vt if pago_bool else 0,
                to_iso_date(data_pagto) if pago_bool and data_pagto else None,
                str(obs).strip() if obs else None,
            ),
        )
        importados += 1
        total += vt

    if importados and not dry_run:
        cur.execute("INSERT INTO historico (quem, acao, detalhe) VALUES (%s, %s, %s)",
                    ("Importação", "Importou planilha",
                     f"{importados} lançamento(s), R$ {total:.2f}, a partir de {desde or 'o início'}"))
    if dry_run:
        conn.rollback()
    else:
        conn.commit()
    conn.close()

    print(("[SIMULACAO, nada foi gravado] " if dry_run else "") + "Resultado:")
    print(f"  Produtos novos ({len(novos)}):", *novos, sep="\n    " if novos else " ")
    print(f"  Produtos com nome/preco alterado ({len(atualizados)}):", *atualizados, sep="\n    " if atualizados else " ")
    print(f"  Lancamentos importados: {importados} (R$ {total:.2f})")
    print(f"  Ja existiam no banco e foram pulados: {repetidos}")
    print(f"  Pessoas novas ({len(pessoas_novas)}):", ", ".join(sorted(set(pessoas_novas))) or "nenhuma")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", nargs="?",
                        default=r"C:\Users\yan.santana\Downloads\PALACIO´S_com_pedido_automatico.xlsx",
                        help="Caminho da planilha")
    parser.add_argument("--reset", action="store_true",
                        help="Apaga tudo que ja existe no banco antes de importar")
    parser.add_argument("--min-row-pedidos", type=int, default=2,
                        help="Primeira linha da aba de lancamentos a importar")
    parser.add_argument("--desde", help="So importa lancamentos desta data em diante (AAAA-MM-DD)")
    parser.add_argument("--dry-run", action="store_true", help="Mostra o que faria, sem gravar")
    parser.add_argument("--categoria-novos", default="mercadinho", choices=["mercadinho", "restaurante"],
                        help="Categoria dos produtos que ainda nao existem no banco")
    args = parser.parse_args()

    try:
        importar(args.xlsx, reset=args.reset, min_row_pedidos=args.min_row_pedidos, desde=args.desde,
                 dry_run=args.dry_run, categoria_novos=args.categoria_novos)
    except FileNotFoundError:
        print(f"Arquivo nao encontrado: {args.xlsx}")
        sys.exit(1)
